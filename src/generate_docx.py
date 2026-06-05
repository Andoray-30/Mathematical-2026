"""
Generate Word document from markdown source.

Converts paper/final_markdown_for_docx.md to paper/final_paper_draft_v2.docx
with proper Chinese academic formatting.

Usage:
    python src/generate_docx.py
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── paths ────────────────────────────────────────────────────────────────
ROOT = Path(__file__).parent.parent
MD_PATH = ROOT / "paper" / "final_markdown_for_docx.md"
OUT_PATH = ROOT / "paper" / "final_paper_draft_v2.docx"
EQUATION_DIR = ROOT / "figures" / "equations"

# ── font names ───────────────────────────────────────────────────────────
FONT_HEI = "SimHei"
FONT_SUN = "SimSun"


# ══════════════════════════════════════════════════════════════════════════
#  Helpers
# ══════════════════════════════════════════════════════════════════════════

def _set_run_font(run, font_name: str, size_pt: float, bold: bool = False):
    """Set font properties on a run, including East-Asian font."""
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = font_name
    # East-Asian font
    rpr = run._element.get_or_add_rPr()
    ea = rpr.find(qn("w:rFonts"))
    if ea is None:
        ea = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{font_name}"/>')
        rpr.insert(0, ea)
    else:
        ea.set(qn("w:eastAsia"), font_name)


def _add_paragraph(doc, text, font_name, size_pt, bold=False,
                   alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   space_before=0, space_after=0,
                   line_spacing=1.25, first_line_indent=None):
    """Add a body paragraph with inline formatting."""
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = Pt(first_line_indent)

    _add_formatted_runs(p, text, font_name, size_pt, bold)
    return p


def _set_run_subscript(run):
    rpr = run._element.get_or_add_rPr()
    for old in rpr.findall(qn("w:vertAlign")):
        rpr.remove(old)
    rpr.append(parse_xml(f'<w:vertAlign {nsdecls("w")} w:val="subScript"/>'))


def _add_math_runs(paragraph, math_text, default_size, default_bold=False):
    subscript_pattern = re.compile(r'([A-Za-z]\w*|\\[A-Za-z]+)_\{([^{}]+)\}')
    pos = 0
    for match in subscript_pattern.finditer(math_text):
        if match.start() > pos:
            run = paragraph.add_run(math_text[pos:match.start()])
            _set_run_font(run, "Cambria Math", default_size * 0.9, bold=default_bold)
            run.font.italic = True

        base_run = paragraph.add_run(match.group(1).lstrip('\\'))
        _set_run_font(base_run, "Cambria Math", default_size * 0.9, bold=default_bold)
        base_run.font.italic = True

        sub_run = paragraph.add_run(match.group(2))
        _set_run_font(sub_run, "Cambria Math", default_size * 0.75, bold=default_bold)
        sub_run.font.italic = True
        _set_run_subscript(sub_run)
        pos = match.end()

    if pos < len(math_text):
        run = paragraph.add_run(math_text[pos:])
        _set_run_font(run, "Cambria Math", default_size * 0.9, bold=default_bold)
        run.font.italic = True


def _add_formatted_runs(paragraph, text, default_font, default_size, default_bold=False):
    """Parse inline markdown (**bold**, *italic*, $math$) and add runs."""
    # Pattern: **bold** | *italic* | $inline_math$ | plain text
    pattern = re.compile(
        r'\*\*(.+?)\*\*'       # bold
        r'|(?<!\*)\*([^*]+?)\*(?!\*)'  # italic (not inside **)
        r'|\$([^$]+?)\$'        # inline math
        r'|([^*$]+)'            # plain text
    )
    for m in pattern.finditer(text):
        if m.group(1) is not None:
            # bold
            run = paragraph.add_run(m.group(1))
            _set_run_font(run, default_font, default_size, bold=True)
        elif m.group(2) is not None:
            # italic
            run = paragraph.add_run(m.group(2))
            _set_run_font(run, default_font, default_size)
            run.font.italic = True
        elif m.group(3) is not None:
            _add_math_runs(paragraph, m.group(3), default_size, default_bold)
        elif m.group(4) is not None:
            # plain text
            plain = m.group(4)
            if plain:
                run = paragraph.add_run(plain)
                _set_run_font(run, default_font, default_size, bold=default_bold)


# ══════════════════════════════════════════════════════════════════════════
#  Equation handling
# ══════════════════════════════════════════════════════════════════════════

def _try_latex_to_omml(latex_str: str):
    """Try to convert LaTeX to OMML XML element. Returns None on failure."""
    try:
        import latex2mathml.converter
        mathml_str = latex2mathml.converter.convert(latex_str)
        # Parse MathML and convert to OMML is non-trivial;
        # python-docx doesn't have a built-in MathML->OMML converter.
        # We'll use a simplified approach: insert as formatted text.
        return None
    except Exception:
        return None


def _add_equation_paragraph(doc, latex_str: str, eq_number: str = ""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25

    image_path = None
    if eq_number:
        image_path = EQUATION_DIR / f"eq_{int(eq_number):02d}.png"

    if image_path and image_path.exists():
        run = p.add_run()
        run.add_picture(str(image_path), width=Cm(12))
    else:
        display = _latex_to_readable(latex_str)
        run = p.add_run(display)
        _set_run_font(run, "Cambria Math", 11, bold=False)
        run.font.italic = True

    if eq_number:
        run2 = p.add_run(f"  ({eq_number})")
        _set_run_font(run2, FONT_SUN, 10.5)

    return p


def _latex_to_readable(latex_str: str) -> str:
    """Convert LaTeX to a readable Unicode-ish string for fallback display."""
    s = latex_str.strip()
    # Common replacements
    replacements = [
        (r'\text{', ""),
        (r'\text', ""),
        (r'\mathrm{', ""),
        (r'\left', ""),
        (r'\right', ""),
        (r'\quad', "  "),
        (r'\qquad', "    "),
        (r'\cdot', "·"),
        (r'\times', "×"),
        (r'\leq', "≤"),
        (r'\geq', "≥"),
        (r'\neq', "≠"),
        (r'\approx', "≈"),
        (r'\infty', "∞"),
        (r'\sum', "Σ"),
        (r'\prod', "Π"),
        (r'\int', "∫"),
        (r'\partial', "∂"),
        (r'\nabla', "∇"),
        (r'\alpha', "α"),
        (r'\beta', "β"),
        (r'\gamma', "γ"),
        (r'\delta', "δ"),
        (r'\epsilon', "ε"),
        (r'\varepsilon', "ε"),
        (r'\theta', "θ"),
        (r'\lambda', "λ"),
        (r'\mu', "μ"),
        (r'\sigma', "σ"),
        (r'\rho', "ρ"),
        (r'\eta', "η"),
        (r'\tilde{x}', "x̃"),
        (r'\max', "max"),
        (r'\min', "min"),
        (r'\med', "med"),
        (r'\text{clip}', "clip"),
        ("\\tag{", "  ("),
        (r'\\', " "),
        (r'\{', "{"),
        (r'\}', "}"),
        (r'\ ', " "),
    ]
    for old, new in replacements:
        s = s.replace(old, new)

    # Clean up remaining \commands
    s = re.sub(r'\\[a-zA-Z]+', '', s)
    # Clean up extra braces
    s = s.replace('{', '').replace('}', '')
    # Clean up multiple spaces
    s = re.sub(r'  +', ' ', s).strip()

    return s


# ══════════════════════════════════════════════════════════════════════════
#  Table handling
# ══════════════════════════════════════════════════════════════════════════

def _parse_table_lines(table_lines: list[str]) -> list[list[str]]:
    """Parse markdown table lines into rows of cells."""
    rows = []
    for line in table_lines:
        line = line.strip()
        if not line.startswith('|'):
            continue
        # Skip separator lines like |---|---|
        if re.match(r'^\|[\s\-:]+\|$', line) or re.match(r'^\|(\s*-+\s*\|)+\s*$', line):
            # Check if it's a separator
            cells_raw = line.split('|')[1:-1]
            if all(re.match(r'^[\s\-:]+$', c) for c in cells_raw):
                continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        rows.append(cells)
    return rows


def _add_three_line_table(doc, rows: list[list[str]], table_title: str = ""):
    """Add a three-line table (三线表) to the document.

    Top border: thick, header separator: thin, bottom border: thick.
    No vertical lines.
    """
    if not rows:
        return

    if table_title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(table_title)
        _set_run_font(run, FONT_HEI, 10.5, bold=True)

    n_cols = max(len(r) for r in rows)
    n_rows = len(rows)

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Calculate column widths - distribute evenly
    page_width_cm = 21.0 - 2.54 * 2  # A4 - margins
    col_width = Cm(page_width_cm / n_cols)

    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j in range(n_cols):
            cell = row.cells[j]
            cell.width = col_width

            # Clear default paragraph
            cell.paragraphs[0].clear()
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            text = row_data[j] if j < len(row_data) else ""
            is_header = (i == 0)
            font_size = 10.0
            _add_formatted_runs(p, text, FONT_SUN, font_size, default_bold=is_header)

    # Apply three-line table borders
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')

    # Table width
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="0" w:type="auto"/>')
    # Remove existing borders
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)

    borders_xml = f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>
        <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>
        <w:left w:val="none" w:sz="0" w:space="0" w:color="000000"/>
        <w:right w:val="none" w:sz="0" w:space="0" w:color="000000"/>
        <w:insideH w:val="none" w:sz="0" w:space="0" w:color="000000"/>
        <w:insideV w:val="none" w:sz="0" w:space="0" w:color="000000"/>
    </w:tblBorders>'''
    tblPr.append(parse_xml(borders_xml))

    # Add thin line under header row
    if n_rows > 1:
        header_row = table.rows[0]
        for cell in header_row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            # Remove existing borders
            for old in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(old)
            cell_borders = f'''<w:tcBorders {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
            </w:tcBorders>'''
            tcPr.append(parse_xml(cell_borders))

    # Add spacing after table
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


# ══════════════════════════════════════════════════════════════════════════
#  Figure placeholder handling
# ══════════════════════════════════════════════════════════════════════════

def _add_figure_placeholder(doc, fig_text: str):
    """Insert a figure image with caption, or a text placeholder if image missing."""
    m = re.match(r'\[插入(图\d+[：:]\s*[^，,]*)[，,]?\s*([^\]]*)\]', fig_text)
    if not m:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(fig_text)
        _set_run_font(run, FONT_SUN, 10.5)
        return p

    label = m.group(1).strip()
    rel_path = m.group(2).strip()
    if label.startswith("图1"):
        rel_path = "figures/final/framework_flowchart_v2.png"
    img_path = ROOT / rel_path if rel_path else None

    if img_path and img_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        run = p_img.add_run()
        run.add_picture(str(img_path), width=Cm(14))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(6)
        cap_run = p_cap.add_run(label)
        _set_run_font(cap_run, FONT_HEI, 10.5, bold=False)
        return p_img
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f"[{label}]")
        _set_run_font(run, FONT_SUN, 10.5, bold=True)
        if rel_path:
            run2 = p.add_run(f"\n({rel_path})")
            _set_run_font(run2, FONT_SUN, 9)
            run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        return p


# ══════════════════════════════════════════════════════════════════════════
#  Markdown parser
# ══════════════════════════════════════════════════════════════════════════

def _classify_line(line: str) -> tuple[str, str]:
    """Classify a markdown line into (block_type, content)."""
    stripped = line.strip()

    if not stripped:
        return ("blank", "")

    # Title (# heading)
    if re.match(r'^#\s+', stripped) and not re.match(r'^##', stripped):
        return ("title", re.sub(r'^#\s+', '', stripped))

    # Heading 1 (##)
    if re.match(r'^##\s+', stripped):
        return ("heading1", re.sub(r'^##\s+', '', stripped))

    # Heading 2 (###)
    if re.match(r'^###\s+', stripped):
        return ("heading2", re.sub(r'^###\s+', '', stripped))

    # Display equation $$...$$
    if stripped.startswith('$$'):
        return ("equation_start", stripped)

    # Figure placeholder
    if stripped.startswith('[插入图') or stripped.startswith('[插图'):
        return ("figure", stripped)

    # Table row
    if stripped.startswith('|') and stripped.endswith('|'):
        return ("table", stripped)

    # Numbered list item (e.g., "1. **text**")
    if re.match(r'^\d+\.\s', stripped):
        return ("list_item", stripped)

    # Bullet list item
    if re.match(r'^[-*]\s', stripped):
        return ("list_item", stripped)

    # Reference line [N] ...
    if re.match(r'^\[\d+\]', stripped):
        return ("reference", stripped)

    # Regular paragraph
    return ("paragraph", stripped)


def parse_markdown(md_text: str) -> list[tuple[str, str]]:
    """Parse markdown into a list of (block_type, content) pairs.

    Handles multi-line equations and tables.
    """
    lines = md_text.split('\n')
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Multi-line equation
        if stripped.startswith('$$'):
            eq_lines = [stripped]
            if stripped.endswith('$$') and len(stripped) > 4:
                # Single-line equation
                blocks.append(("equation", stripped))
                i += 1
                continue
            # Multi-line: collect until closing $$
            i += 1
            while i < n:
                eq_lines.append(lines[i].strip())
                if lines[i].strip().endswith('$$'):
                    break
                i += 1
            i += 1
            blocks.append(("equation", ' '.join(eq_lines)))
            continue

        # Table: collect consecutive table lines
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = [stripped]
            i += 1
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            blocks.append(("table_block", table_lines))
            continue

        # Regular line
        btype, content = _classify_line(line)
        if btype != "blank":  # Skip blank lines (they're separators)
            blocks.append((btype, content))
        i += 1

    return blocks


# ══════════════════════════════════════════════════════════════════════════
#  Document builder
# ══════════════════════════════════════════════════════════════════════════

def _extract_eq_number(eq_text: str) -> tuple[str, str]:
    """Extract equation number from \\tag{N} in equation text."""
    m = re.search(r'\\tag\{(\d+)\}', eq_text)
    if m:
        eq_num = m.group(1)
        clean = eq_text.replace(m.group(0), '').strip()
        # Remove $$ wrappers
        clean = re.sub(r'^\$\$\s*', '', clean)
        clean = re.sub(r'\s*\$\$$', '', clean)
        return clean, eq_num
    # Remove $$ wrappers
    clean = re.sub(r'^\$\$\s*', '', eq_text)
    clean = re.sub(r'\s*\$\$$', '', clean)
    return clean, ""


def build_document(blocks: list[tuple[str, str]]) -> Document:
    """Build a Word document from parsed blocks."""
    doc = Document()

    # ── Page setup ────────────────────────────────────────────────────────
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    # ── Page numbers in footer ────────────────────────────────────────────
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    _set_run_font(run, FONT_SUN, 9)
    # PAGE field
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._element.append(fldChar1)
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    run._element.append(instrText)
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._element.append(fldChar2)

    # ── No header ─────────────────────────────────────────────────────────
    header = section.header
    header.is_linked_to_previous = False

    # ── Content ───────────────────────────────────────────────────────────
    # State for tracking section context
    in_references = False
    table_title_pending = ""

    for block_type, content in blocks:

        # Title
        if block_type == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(content)
            _set_run_font(run, FONT_HEI, 16, bold=True)
            continue

        # Heading 1 (##)
        if block_type == "heading1":
            # Track if we're in references section
            if "参考文献" in content:
                in_references = True
            else:
                in_references = False

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(content)
            _set_run_font(run, FONT_HEI, 14, bold=True)
            continue

        # Heading 2 (###)
        if block_type == "heading2":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(content)
            _set_run_font(run, FONT_HEI, 12, bold=True)
            continue

        # Equation
        if block_type == "equation":
            eq_text, eq_num = _extract_eq_number(content)
            _add_equation_paragraph(doc, eq_text, eq_num)
            continue

        # Figure placeholder
        if block_type == "figure":
            _add_figure_placeholder(doc, content)
            continue

        # Table block
        if block_type == "table_block":
            # Check if previous block was a table title paragraph
            rows = _parse_table_lines(content)
            title = table_title_pending
            table_title_pending = ""
            _add_three_line_table(doc, rows, title)
            continue

        # Reference
        if block_type == "reference":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            # Hanging indent for references
            p.paragraph_format.first_line_indent = Pt(-21)
            p.paragraph_format.left_indent = Pt(21)
            _add_formatted_runs(p, content, FONT_SUN, 9)
            continue

        # List item
        if block_type == "list_item":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.first_line_indent = Pt(21)
            _add_formatted_runs(p, content, FONT_SUN, 10.5)
            continue

        # Regular paragraph
        if block_type == "paragraph":
            if (content.startswith("队伍编号：") or
                    content.startswith("选题：") or
                    content.startswith("论文题目：")):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(content)
                _set_run_font(run, FONT_HEI, 16, bold=False)
                continue

            if content.startswith("**关键词**") or content.startswith("关键词"):
                _add_paragraph(
                    doc, content, FONT_SUN, 10.5,
                    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                    space_before=3,
                    space_after=3,
                    line_spacing=1.25,
                    first_line_indent=21,
                )
                doc.add_page_break()
                continue

            # Check if this looks like a table title (e.g., "**表1 ...**")
            if re.match(r'^\*\*表\d+', content):
                table_title_pending = content.replace('**', '')
                continue

            # Detect AI tool statement and appendix header
            is_statement = "AI工具使用声明" in content or "AI工具使用" in content
            is_appendix_header = content.startswith("附录") and len(content) < 20

            font_size = 10.5
            space_before = 3
            space_after = 3

            if in_references:
                # Shouldn't reach here (handled above), but just in case
                font_size = 9

            p = _add_paragraph(
                doc, content, FONT_SUN, font_size,
                bold=is_statement,
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                space_before=space_before,
                space_after=space_after,
                line_spacing=1.25,
                first_line_indent=21 if not is_statement and not is_appendix_header else None,
            )

            if is_statement:
                p.paragraph_format.space_before = Pt(8)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            continue

    return doc


# ══════════════════════════════════════════════════════════════════════════
#  Main
# ══════════════════════════════════════════════════════════════════════════

def main() -> int:
    if not MD_PATH.exists():
        print(f"ERROR: Source markdown not found: {MD_PATH}")
        return 1

    print(f"Reading: {MD_PATH}")
    md_text = MD_PATH.read_text(encoding="utf-8")
    print(f"  Lines: {len(md_text.splitlines())}")

    print("Parsing markdown...")
    blocks = parse_markdown(md_text)
    print(f"  Blocks: {len(blocks)}")

    print("Building document...")
    doc = build_document(blocks)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    size_kb = OUT_PATH.stat().st_size / 1024
    print(f"Saved: {OUT_PATH}  ({size_kb:.1f} KB)")

    if size_kb < 50:
        print("WARNING: Output file is smaller than expected (< 50 KB)")
    else:
        print("OK: File size > 50 KB")

    return 0


if __name__ == "__main__":
    sys.exit(main())
